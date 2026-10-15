from __future__ import annotations

import argparse
import importlib.util
import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(r"D:\GitHub\AppSheet Project")
WORKBOOK = ROOT / "Documents" / "AppSheet_Data_Dictionary.xlsx"
OUTPUT_DIR = ROOT / "outputs" / "olap-daily-installment-2026-09-01"
OUTPUT_WORKBOOK = OUTPUT_DIR / "AppSheet_Data_Dictionary.xlsx"
RENDER_HELPER = ROOT / "outputs" / "borrower-visibility-2026-09-01" / "update_dictionary.py"


def find_row(ws, table: str, column: str, start: int = 5) -> int:
    for row in range(start, ws.max_row + 1):
        if ws.cell(row, 1).value == table and ws.cell(row, 3).value == column:
            return row
    raise ValueError(f"Could not find {table}[{column}] in {ws.title}")


def find_value_row(ws, column: int, value: str, start: int = 1) -> int:
    for row in range(start, ws.max_row + 1):
        if ws.cell(row, column).value == value:
            return row
    raise ValueError(f"Could not find {value!r} in {ws.title} column {column}")


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def set_row(ws, row: int, values: list) -> None:
    for col, value in enumerate(values, start=1):
        ws.cell(row, col).value = value


def extract_oltp_formula(value: str) -> str:
    if isinstance(value, str) and value.startswith("OLTP: ") and "\nOLAP: " in value:
        return value[len("OLTP: ") : value.index("\nOLAP: ")]
    return value


PROJECTED_OUTSTANDING = '''SUM(
  SELECT(
    Loans[Outstanding Principal],
    AND(
      [Ref Borrowers] = [_THISROW].[Assessment Borrower],
      [Loan Status] = "ยังไม่ปิดยอด",
      OR(
        [Loan Type] = "ดอกเบี้ยรายวัน",
        AND([Loan Type] = "กำหนดวันชำระ", [Due Date] > [_THISROW].[Date])
      )
    )
  )
)
+ MAX(
  LIST(
    [Assessment Installment Future Principal Due] * 0,
    SUM(
      SELECT(
        Loans[Outstanding Principal],
        AND(
          [Ref Borrowers] = [_THISROW].[Assessment Borrower],
          [Loan Status] = "ยังไม่ปิดยอด",
          [Loan Type] = "ผ่อนชำระรายวัน",
          [Loan Date] <= [_THISROW].[Date]
        )
      )
    )
    - SUM(
      SELECT(
        Calendar[Assessment Installment Future Principal Due],
        [Date] <= [_THISROW].[Date]
      )
    )
  )
)'''

PROJECTED_INTEREST = '''SUM(
  SELECT(
    Loans[Total Interest Received],
    AND(
      [Ref Borrowers] = [_THISROW].[Assessment Borrower],
      [Loan Status] = "ยังไม่ปิดยอด"
    )
  )
)
+ (
  SUM(
    SELECT(
      Loans[Current Daily Interest],
      AND(
        [Ref Borrowers] = [_THISROW].[Assessment Borrower],
        [Loan Status] = "ยังไม่ปิดยอด",
        [Loan Type] = "ดอกเบี้ยรายวัน"
      )
    )
  )
  * IF([Date] > TODAY(), TOTALHOURS([Date] - TODAY()) / 24, 0)
)
+ SUM(
  SELECT(
    Loans[Fixed Interest],
    AND(
      [Ref Borrowers] = [_THISROW].[Assessment Borrower],
      [Loan Status] = "ยังไม่ปิดยอด",
      [Loan Type] = "กำหนดวันชำระ",
      [Due Date] > TODAY(),
      [Due Date] <= [_THISROW].[Date]
    )
  )
)
+ SUM(
  SELECT(
    Calendar[Assessment Installment Future Interest Due],
    [Date] <= [_THISROW].[Date]
  )
)'''

ASSESSMENT_BASE = '''SUM(
  SELECT(
    Loans[Principal Amount],
    AND(
      [Ref Borrowers] = [_THISROW].[Assessment Borrower],
      [Loan Status] = "ยังไม่ปิดยอด",
      OR(
        [Loan Type] = "ดอกเบี้ยรายวัน",
        AND([Loan Type] = "กำหนดวันชำระ", [Due Date] > [_THISROW].[Date])
      )
    )
  )
)
+ [Assessment Installment Contract Outstanding]'''

ASSESSMENT_WEIGHT = '''SUM(
  SELECT(
    Loans[Principal Start Weight],
    AND(
      [Ref Borrowers] = [_THISROW].[Assessment Borrower],
      [Loan Status] = "ยังไม่ปิดยอด",
      OR(
        [Loan Type] = "ดอกเบี้ยรายวัน",
        AND([Loan Type] = "กำหนดวันชำระ", [Due Date] > [_THISROW].[Date])
      )
    )
  )
)
+ (([Candidate Day Index] + 1) * [Assessment Installment Contract Outstanding])
- SUM(
  SELECT(
    Calendar[Assessment Installment Contract Outstanding],
    [Date] <= [_THISROW].[Date]
  )
)'''

BASE_PRINCIPAL = '''IF(
  AND(
    [Loan Type] = "ผ่อนชำระรายวัน",
    ISNOTBLANK([Loan Date]),
    ISNOTBLANK([Due Date]),
    [Due Date] >= [Loan Date]
  ),
  FLOOR(
    [Principal Amount]
    / (TOTALHOURS([Due Date] - [Loan Date]) / 24 + 1)
  ),
  0
)'''

CONTRACT_PRINCIPAL_DUE = '''SUM(
  SELECT(
    Loans[Daily Installment Base Principal],
    AND(
      [Ref Borrowers] = [_THISROW].[Assessment Borrower],
      [Loan Status] = "ยังไม่ปิดยอด",
      [Loan Type] = "ผ่อนชำระรายวัน",
      [Loan Date] <= [_THISROW].[Date],
      [Due Date] >= [_THISROW].[Date],
      [Due Date] >= [Loan Date],
      [Daily Payment Amount] > 0
    )
  )
)
+ COUNT(
  SELECT(
    Loans[Row ID],
    AND(
      [Ref Borrowers] = [_THISROW].[Assessment Borrower],
      [Loan Status] = "ยังไม่ปิดยอด",
      [Loan Type] = "ผ่อนชำระรายวัน",
      [Loan Date] <= [_THISROW].[Date],
      [Due Date] >= [_THISROW].[Date],
      [Due Date] >= [Loan Date],
      [Daily Payment Amount] > 0,
      (TOTALHOURS([_THISROW].[Date] - [Loan Date]) / 24 + 1)
      <= MOD(
        [Principal Amount],
        (TOTALHOURS([Due Date] - [Loan Date]) / 24 + 1)
      )
    )
  )
)'''

FUTURE_PRINCIPAL_DUE = CONTRACT_PRINCIPAL_DUE.replace(
    "[Daily Payment Amount] > 0\n    )",
    '''[Daily Payment Amount] > 0,
      [_THISROW].[Date] > IF(
        ISNOTBLANK(MAX([Related Charges][Charge Date])),
        MAX([Related Charges][Charge Date]),
        [Loan Date] - 1
      )
    )''',
).replace(
    "[Daily Payment Amount] > 0,\n      (TOTALHOURS",
    '''[Daily Payment Amount] > 0,
      [_THISROW].[Date] > IF(
        ISNOTBLANK(MAX([Related Charges][Charge Date])),
        MAX([Related Charges][Charge Date]),
        [Loan Date] - 1
      ),
      (TOTALHOURS''',
)

FUTURE_INTEREST_DUE = '''SUM(
  SELECT(
    Loans[Daily Payment Amount],
    AND(
      [Ref Borrowers] = [_THISROW].[Assessment Borrower],
      [Loan Status] = "ยังไม่ปิดยอด",
      [Loan Type] = "ผ่อนชำระรายวัน",
      [Loan Date] <= [_THISROW].[Date],
      [Due Date] >= [_THISROW].[Date],
      [Due Date] >= [Loan Date],
      [Daily Payment Amount] > 0,
      [_THISROW].[Date] > IF(
        ISNOTBLANK(MAX([Related Charges][Charge Date])),
        MAX([Related Charges][Charge Date]),
        [Loan Date] - 1
      )
    )
  )
)
- [Assessment Installment Future Principal Due]'''

CONTRACT_OUTSTANDING = '''MAX(
  LIST(
    [Assessment Installment Contract Principal Due] * 0,
    SUM(
      SELECT(
        Loans[Principal Amount],
        AND(
          [Ref Borrowers] = [_THISROW].[Assessment Borrower],
          [Loan Status] = "ยังไม่ปิดยอด",
          [Loan Type] = "ผ่อนชำระรายวัน",
          [Loan Date] <= [_THISROW].[Date],
          [Due Date] >= [Loan Date],
          [Daily Payment Amount] > 0
        )
      )
    )
    - SUM(
      SELECT(
        Calendar[Assessment Installment Contract Principal Due],
        [Date] <= [_THISROW].[Date]
      )
    )
  )
)'''


def update_workbook() -> None:
    wb = load_workbook(WORKBOOK)

    readme = wb["README"]
    readme.cell(find_value_row(readme, 1, "Reconciliation status"), 2).value = (
        "Reconciled with production OLTP v2.200099 and OLAP v1.000164 on 2026-09-01. "
        "OLAP now loads Loans[Daily Payment Amount], calculates daily-installment expected interest, "
        "uses declining whole-baht installment schedules in Loan Assessment, and exposes both governed "
        "daily-installment fields through read-only Reporting Loans. Borrowers[Hidden Flag] remains "
        "intentionally OLTP-only and is the sole live schema-drift warning in OLAP."
    )
    readme.cell(find_value_row(readme, 1, "Live baselines"), 2).value = (
        "OLTP v2.200099 — Loan Manager ผู้จัดการสินเชื่อ — 93844270-6aa5-420b-a960-bcdf08abd357\n"
        "OLAP v1.000164 — Loan Business Insights — e74dd368-654f-44d7-bccf-83cf361d1cdb"
    )
    readme.cell(24, 1).value = (
        "Shared VC baseline: 82 VCs are present in both apps and the five daily-installment formula "
        "variances are resolved. Loans[Daily Installment Base Principal] and the four Calendar schedule "
        "VCs are OLAP-only reporting helpers. Borrowers[Hidden Flag] remains OLTP-only."
    )

    tables = wb["Tables"]
    calendar_table = find_value_row(tables, 1, "Calendar", start=5)
    tables.cell(calendar_table, 6).value = 26
    tables.cell(calendar_table, 7).value = (
        "Google Sheets-backed date spine used by OLAP loan-assessment projections, including "
        "declining daily-installment principal and interest schedules."
    )
    loans_table = find_value_row(tables, 1, "Loans", start=5)
    tables.cell(loans_table, 6).value = 30
    tables.cell(loans_table, 7).value = (
        "Loan contracts, commercial terms, lifecycle state, balances, daily-installment reporting, "
        "and loan-level risk metrics."
    )

    columns = wb["Columns"]

    daily_payment = find_row(columns, "Loans", "Daily Payment Amount")
    columns.cell(daily_payment, 5).value = "OLTP + OLAP"
    columns.cell(daily_payment, 6).value = (
        'OLTP: Required_If/Show_If and whole-baht validation for "ผ่อนชำระรายวัน"; '
        "OLAP: loaded read-only; decimal digits 0"
    )
    columns.cell(daily_payment, 8).value = (
        'Stored source column; display name: "ยอดชำระรายวัน / Daily Payment Amount"; '
        "OLTP form validation remains app-specific; OLAP description states that interval charges combine elapsed days"
    )

    for table, column in [
        ("Loans", "Expected Interest This Month"),
        ("Loans", "Effective Daily Interest Rate"),
        ("Loans", "Loan Type TE"),
        ("Borrowers", "Active Daily Interest"),
        ("Statistics", "Expected Daily Interest"),
    ]:
        row = find_row(columns, table, column)
        columns.cell(row, 6).value = "-"
        columns.cell(row, 8).value = extract_oltp_formula(columns.cell(row, 8).value)

    expected_month = find_row(columns, "Loans", "Expected Interest This Month")
    columns.cell(expected_month, 7).value = (
        "Open-loan current-month interest forecast. Daily-installment interest covers only uncharged "
        "contract days after the latest posted Charge through the earlier of Due Date and month-end."
    )
    expected_amount = find_row(columns, "Loans", "Expected Daily Interest Amount")
    columns.cell(expected_amount, 5).value = "OLTP + OLAP"
    columns.cell(expected_amount, 6).value = "Hidden in OLTP + OLAP; OLAP Price with 2 decimal digits"

    # Append the one OLAP-only loan helper at the end of the Loans block.
    loan_helper_row = expected_amount + 1
    columns.insert_rows(loan_helper_row, 1)
    copy_row_style(columns, expected_amount, loan_helper_row)
    set_row(
        columns,
        loan_helper_row,
        [
            "Loans",
            50,
            "Daily Installment Base Principal",
            "VC / Number",
            "OLAP only",
            "Hidden; whole-baht base allocation",
            "Base principal allocated to every contract day before the one-baht earliest-day remainder is applied.",
            BASE_PRINCIPAL,
        ],
    )

    # Update the four existing Calendar formulas.
    calendar_updates = {
        "Projected Outstanding Principal": (
            "Projected selected-borrower principal. Daily-installment balances start from current Outstanding Principal, "
            "subtract only contract principal after the latest posted Charge through the candidate date/Due Date, and do not go below zero.",
            PROJECTED_OUTSTANDING,
        ),
        "Projected Interest": (
            "Actual selected-borrower interest received plus future daily-interest, fixed-date, and uncharged daily-installment interest. "
            "Daily-installment Transfer Fee is not projected again.",
            PROJECTED_INTEREST,
        ),
        "Assessment Principal Base": (
            "Principal base used by Minimum Required Profit; daily-installment exposure follows the declining contract balance.",
            ASSESSMENT_BASE,
        ),
        "Assessment Principal Start Weight": (
            "Principal-day start weight used by Minimum Required Profit; algebraically preserves cumulative declining daily-installment principal-days.",
            ASSESSMENT_WEIGHT,
        ),
    }
    for name, (purpose, formula) in calendar_updates.items():
        row = find_row(columns, "Calendar", name)
        columns.cell(row, 7).value = purpose
        columns.cell(row, 8).value = formula

    calendar_rows = [
        (
            "Assessment Installment Contract Principal Due",
            "Whole-baht scheduled principal for the candidate Calendar date; FLOOR base plus one baht on the earliest remainder days.",
            CONTRACT_PRINCIPAL_DUE,
        ),
        (
            "Assessment Installment Future Principal Due",
            "Whole-baht scheduled principal after each loan's latest posted Charge through its Due Date.",
            FUTURE_PRINCIPAL_DUE,
        ),
        (
            "Assessment Installment Future Interest Due",
            "Scheduled daily-installment payment after the latest Charge minus future scheduled principal; Transfer Fee excluded.",
            FUTURE_INTEREST_DUE,
        ),
        (
            "Assessment Installment Contract Outstanding",
            "Nonnegative deterministic daily-installment contract balance after cumulative scheduled principal through the Calendar date.",
            CONTRACT_OUTSTANDING,
        ),
    ]
    calendar_last = max(
        row for row in range(5, columns.max_row + 1) if columns.cell(row, 1).value == "Calendar"
    )
    insert_at = calendar_last + 1
    columns.insert_rows(insert_at, len(calendar_rows))
    for offset, (name, purpose, formula) in enumerate(calendar_rows):
        target = insert_at + offset
        copy_row_style(columns, calendar_last, target)
        set_row(
            columns,
            target,
            [
                "Calendar",
                25 + offset,
                name,
                "VC / Price",
                "OLAP only",
                "Hidden; whole baht",
                purpose,
                formula,
            ],
        )

    columns.tables["DataModelColumns"].ref = f"A4:H{columns.max_row}"

    slices = wb["Slices"]
    reporting_loans = next(
        row
        for row in range(5, slices.max_row + 1)
        if slices.cell(row, 1).value == "OLAP" and slices.cell(row, 2).value == "Reporting Loans"
    )
    slices.cell(reporting_loans, 6).value = 32
    slices.cell(reporting_loans, 7).value = (
        "Curated, read-only loan feed containing contract, balance, daily-installment amount/expected interest, rate, age, and default fields."
    )

    slice_columns = wb["Slice Columns"]
    reporting_rows = [
        row
        for row in range(5, slice_columns.max_row + 1)
        if slice_columns.cell(row, 1).value == "OLAP"
        and slice_columns.cell(row, 2).value == "Reporting Loans"
    ]
    append_at = max(reporting_rows) + 1
    slice_columns.insert_rows(append_at, 2)
    for offset, name in enumerate(["Daily Payment Amount", "Expected Daily Interest Amount"]):
        target = append_at + offset
        copy_row_style(slice_columns, append_at - 1, target)
        set_row(slice_columns, target, ["OLAP", "Reporting Loans", 31 + offset, name])
    slice_columns.tables["DataModelSliceColumns"].ref = f"A4:D{slice_columns.max_row}"

    variances = wb["App Variances"]
    delete_rows = [
        row
        for row in range(5, variances.max_row + 1)
        if variances.cell(row, 1).value == "Column"
        and variances.cell(row, 2).value == "Loans"
        and variances.cell(row, 3).value in {"Daily Payment Amount", "Expected Daily Interest Amount"}
    ]
    for row in sorted(delete_rows, reverse=True):
        variances.delete_rows(row, 1)

    calendar_variance_last = max(
        row for row in range(5, variances.max_row + 1) if variances.cell(row, 2).value == "Calendar"
    )
    variances.insert_rows(calendar_variance_last + 1, len(calendar_rows))
    for offset, (name, purpose, _formula) in enumerate(calendar_rows):
        target = calendar_variance_last + 1 + offset
        copy_row_style(variances, calendar_variance_last, target)
        set_row(variances, target, ["Column", "Calendar", name, "No", "Yes", purpose])

    loans_variance_last = max(
        row for row in range(5, variances.max_row + 1) if variances.cell(row, 2).value == "Loans"
    )
    variances.insert_rows(loans_variance_last + 1, 1)
    copy_row_style(variances, loans_variance_last, loans_variance_last + 1)
    set_row(
        variances,
        loans_variance_last + 1,
        [
            "Column",
            "Loans",
            "Daily Installment Base Principal",
            "No",
            "Yes",
            "Hidden whole-baht base allocation used only by OLAP Calendar projection helpers.",
        ],
    )
    variances.tables["DataModelAppVariances"].ref = f"A4:F{variances.max_row}"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(WORKBOOK)
    shutil.copy2(WORKBOOK, OUTPUT_WORKBOOK)


def render(phase: str) -> None:
    spec = importlib.util.spec_from_file_location("render_helper", RENDER_HELPER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    module.WORKBOOK = WORKBOOK
    module.OUTPUT_DIR = OUTPUT_DIR
    module.render_workbook(phase)


def render_focus() -> None:
    spec = importlib.util.spec_from_file_location("render_helper", RENDER_HELPER)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    source = load_workbook(WORKBOOK, data_only=False)
    focus_dir = OUTPUT_DIR / "render-focus"
    focus_dir.mkdir(parents=True, exist_ok=True)

    selections = {
        "Controls": (
            "README",
            list(range(1, 25)),
        ),
        "Tables": (
            "Tables",
            [1, 2, 3, 4]
            + [
                row
                for row in range(5, source["Tables"].max_row + 1)
                if source["Tables"].cell(row, 1).value in {"Calendar", "Loans"}
            ],
        ),
        "Columns Focus": (
            "Columns",
            [1, 2, 3, 4]
            + [
                row
                for row in range(5, source["Columns"].max_row + 1)
                if (source["Columns"].cell(row, 1).value, source["Columns"].cell(row, 3).value)
                in {
                    ("Borrowers", "Active Daily Interest"),
                    ("Statistics", "Expected Daily Interest"),
                    ("Loans", "Daily Payment Amount"),
                    ("Loans", "Expected Daily Interest Amount"),
                    ("Loans", "Daily Installment Base Principal"),
                    ("Loans", "Expected Interest This Month"),
                    ("Loans", "Effective Daily Interest Rate"),
                    ("Loans", "Loan Type TE"),
                    ("Calendar", "Projected Outstanding Principal"),
                    ("Calendar", "Projected Interest"),
                    ("Calendar", "Assessment Principal Base"),
                    ("Calendar", "Assessment Principal Start Weight"),
                    ("Calendar", "Assessment Installment Contract Principal Due"),
                    ("Calendar", "Assessment Installment Future Principal Due"),
                    ("Calendar", "Assessment Installment Future Interest Due"),
                    ("Calendar", "Assessment Installment Contract Outstanding"),
                }
            ],
        ),
        "Slices Focus": (
            "Slices",
            [1, 2, 3, 4]
            + [
                row
                for row in range(5, source["Slices"].max_row + 1)
                if source["Slices"].cell(row, 2).value == "Reporting Loans"
            ],
        ),
        "Slice Columns Focus": (
            "Slice Columns",
            [1, 2, 3, 4]
            + [
                row
                for row in range(5, source["Slice Columns"].max_row + 1)
                if source["Slice Columns"].cell(row, 2).value == "Reporting Loans"
                and source["Slice Columns"].cell(row, 3).value in {31, 32}
            ],
        ),
        "Variances Focus": (
            "App Variances",
            [1, 2, 3, 4]
            + [
                row
                for row in range(5, source["App Variances"].max_row + 1)
                if (source["App Variances"].cell(row, 1).value, source["App Variances"].cell(row, 2).value, source["App Variances"].cell(row, 3).value)
                in {
                    ("Column", "Loans", "Daily Installment Base Principal"),
                    ("Column", "Calendar", "Assessment Installment Contract Principal Due"),
                    ("Column", "Calendar", "Assessment Installment Future Principal Due"),
                    ("Column", "Calendar", "Assessment Installment Future Interest Due"),
                    ("Column", "Calendar", "Assessment Installment Contract Outstanding"),
                    ("Column", "Borrowers", "Hidden Flag"),
                }
            ],
        ),
    }

    previews = []
    for index, (title, (sheet_name, rows)) in enumerate(selections.items(), start=1):
        ws = source.create_sheet(title)
        source_ws = source[sheet_name]
        for col_letter, dimension in source_ws.column_dimensions.items():
            ws.column_dimensions[col_letter].width = dimension.width
        for target_row, source_row in enumerate(rows, start=1):
            ws.row_dimensions[target_row].height = source_ws.row_dimensions[source_row].height
            for col in range(1, source_ws.max_column + 1):
                src = source_ws.cell(source_row, col)
                dst = ws.cell(target_row, col, src.value)
                if src.has_style:
                    dst._style = copy(src._style)
                dst.number_format = src.number_format
                dst.alignment = copy(src.alignment)
        path = focus_dir / f"{index:02d}-{title.replace(' ', '_')}.png"
        module.render_sheet(ws, path)
        previews.append(path)
    print({"focus_previews": [str(path) for path in previews]})


def validate() -> None:
    wb = load_workbook(WORKBOOK, data_only=False)
    assert len(wb.sheetnames) == 8
    columns = wb["Columns"]
    slices = wb["Slices"]
    slice_columns = wb["Slice Columns"]
    variances = wb["App Variances"]

    loans = [row for row in columns.iter_rows(min_row=5, values_only=True) if row[0] == "Loans"]
    calendar = [row for row in columns.iter_rows(min_row=5, values_only=True) if row[0] == "Calendar"]
    assert len(loans) == 50
    assert len(calendar) == 28
    assert len({row[2] for row in loans}) == 50
    assert len({row[2] for row in calendar}) == 28

    by_key = {(row[0], row[2]): row for row in columns.iter_rows(min_row=5, values_only=True)}
    assert by_key[("Loans", "Daily Payment Amount")][4] == "OLTP + OLAP"
    assert by_key[("Loans", "Expected Daily Interest Amount")][4] == "OLTP + OLAP"
    assert by_key[("Loans", "Daily Installment Base Principal")][4] == "OLAP only"
    for key in [
        ("Loans", "Expected Interest This Month"),
        ("Loans", "Effective Daily Interest Rate"),
        ("Loans", "Loan Type TE"),
        ("Borrowers", "Active Daily Interest"),
        ("Statistics", "Expected Daily Interest"),
    ]:
        assert by_key[key][5] == "-"
        assert "OLAP:" not in str(by_key[key][7])

    assert len([row for row in loans if row[2] == "Daily Installment Base Principal"]) == 1
    calendar_helper_names = [
        "Assessment Installment Contract Principal Due",
        "Assessment Installment Future Principal Due",
        "Assessment Installment Future Interest Due",
        "Assessment Installment Contract Outstanding",
    ]
    for name in calendar_helper_names:
        assert len([row for row in calendar if row[2] == name]) == 1

    reporting_slice = [
        row
        for row in slices.iter_rows(min_row=5, values_only=True)
        if row[0] == "OLAP" and row[1] == "Reporting Loans"
    ]
    assert len(reporting_slice) == 1 and reporting_slice[0][5] == 32
    reporting_members = [
        row
        for row in slice_columns.iter_rows(min_row=5, values_only=True)
        if row[0] == "OLAP" and row[1] == "Reporting Loans"
    ]
    assert len(reporting_members) == 32
    assert reporting_members[-2][2:] == (31, "Daily Payment Amount")
    assert reporting_members[-1][2:] == (32, "Expected Daily Interest Amount")

    variance_keys = {
        (row[1], row[2]) for row in variances.iter_rows(min_row=5, values_only=True)
    }
    assert ("Loans", "Daily Payment Amount") not in variance_keys
    assert ("Loans", "Expected Daily Interest Amount") not in variance_keys
    assert ("Loans", "Daily Installment Base Principal") in variance_keys
    for name in calendar_helper_names:
        assert ("Calendar", name) in variance_keys

    assert columns.tables["DataModelColumns"].ref == f"A4:H{columns.max_row}"
    assert slice_columns.tables["DataModelSliceColumns"].ref == f"A4:D{slice_columns.max_row}"
    assert variances.tables["DataModelAppVariances"].ref == f"A4:F{variances.max_row}"
    assert wb["Tables"].cell(find_value_row(wb["Tables"], 1, "Loans", 5), 6).value == 30
    assert wb["Tables"].cell(find_value_row(wb["Tables"], 1, "Calendar", 5), 6).value == 26

    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")
    errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(token in cell.value for token in error_tokens):
                    errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    assert not errors, errors

    print(
        {
            "workbook": str(WORKBOOK),
            "output_copy": str(OUTPUT_WORKBOOK),
            "sheets": len(wb.sheetnames),
            "columns": columns.max_row - 4,
            "loans_columns": len(loans),
            "calendar_columns": len(calendar),
            "reporting_loans_columns": len(reporting_members),
            "app_variances": variances.max_row - 4,
            "formula_errors": errors,
        }
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["update", "render-after", "render-focus", "validate"])
    args = parser.parse_args()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if args.mode == "update":
        update_workbook()
    elif args.mode == "render-after":
        render("after")
    elif args.mode == "render-focus":
        render_focus()
    else:
        validate()


if __name__ == "__main__":
    main()
