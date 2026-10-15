import json
from pathlib import Path
from collections import Counter
import openpyxl

p=Path(__file__).parent
a=openpyxl.load_workbook(p/'AppSheet_Data_Dictionary.before.xlsx')
b=openpyxl.load_workbook(p/'AppSheet_Data_Dictionary.xlsx')
assert a.sheetnames==b.sheetnames
def rows(w,s,n):
    return [tuple(r[:n]) for r in w[s].iter_rows(min_row=5,values_only=True) if r[0]]
old=rows(a,'Columns',8); new=rows(b,'Columns',8)
assert len(new)==365 and len(new)-len(old)==47
assert len({(r[0],r[2]) for r in new})==len(new)
lookup={(r[0],r[2]):r for r in new}
for r in old:
    if r[4]=='OLAP only': assert lookup[r[0],r[2]]==r
assert Counter(r[0] for r in new)['Payments']==23
assert Counter(r[0] for r in new)['Payment Allocations']==19
assert len(rows(b,'Relationships',5))-len(rows(a,'Relationships',5))==6
assert len(rows(b,'Tables',7))==13
for s,n in [('Slices',8),('Slice Columns',4)]:
    assert [r for r in rows(a,s,n) if r[0]=='OLAP']==[r for r in rows(b,s,n) if r[0]=='OLAP']
sc=[r for r in rows(b,'Slice Columns',4) if r[0]=='OLTP']
assert len(sc)==45
expected=json.loads((p.parent/'universal-payments-20260905/clone_slices_checkpoint065.json').read_text(encoding='utf-8'))
for s in expected:
    assert [r[3] for r in sc if r[1]==s['name']]==s['columns']
for s in ['Payments Clone','Payments Clone Slices','Change Log']:
    assert list(a[s].values)==list(b[s].values)
assert not [(s.title,c.coordinate) for s in b for row in s for c in row if c.data_type=='e']
result={'status':'PASS','sheets_preserved':10,'union_columns':365,'new_union_columns':47,'new_relationships':6,'oltp_slice_membership_rows':45,'olap_only_column_definitions_and_slice_memberships':'unchanged','historical_clone_sheets':'unchanged','excel_error_cells':0}
(p/'dictionary_verification.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result))
