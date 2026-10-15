"""Fallback narrow XLSX edit when bundled artifact-tool is unavailable.

Only the existing Manifest D2/E2 XML value nodes change. Other ZIP members and
all workbook features remain byte-identical. openpyxl is read-only validation.
"""
import argparse,hashlib,json,re,textwrap
from pathlib import Path
from zipfile import ZipFile
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from PIL import Image,ImageDraw,ImageFont

p=argparse.ArgumentParser();p.add_argument('source');p.add_argument('output');p.add_argument('--edit',action='store_true');a=p.parse_args()
source=Path(a.source);out=Path(a.output);out.mkdir(parents=True,exist_ok=True)
w=load_workbook(source);s=w['Manifest']
assert [s['D2'].value,s['E2'].value]==['Loan Business Insights','e74dd368-654f-44d7-bccf-83cf361d1cdb']
changes={'D2':'<<_APPNAME>>','E2':'<<_APPID>>'}
def render(sheet,path):
    # Limited cell preview of unchanged column widths/row heights and formatting.
    cols=['C','D','E','F','G'];scale=2
    widths=[round((sheet.column_dimensions[col].width*7+5)*scale) for col in cols]
    heights=[round((sheet.row_dimensions[row].height or 15)*4/3*scale) for row in [1,2]]
    im=Image.new('RGB',(sum(widths),sum(heights)),'white');d=ImageDraw.Draw(im);y=0
    for row,h in zip([1,2],heights):
        x=0
        for col,width in zip(cols,widths):
            c=sheet[f'{col}{row}'];fill=c.fill.fgColor
            if fill.type=='rgb' and c.fill.patternType=='solid':d.rectangle((x,y,x+width,y+h),fill='#'+fill.rgb[-6:])
            font=ImageFont.truetype('C:/Windows/Fonts/calibrib.ttf' if c.font.b else 'C:/Windows/Fonts/calibri.ttf',round(c.font.sz*4/3*scale))
            color='#'+c.font.color.rgb[-6:] if c.font.color and c.font.color.type=='rgb' else '#000000'
            lines=[];line=''
            for ch in str(c.value or ''):
                if d.textlength(line+ch,font=font)>width-10:lines.append(line);line=ch
                else:line+=ch
            lines.append(line);d.multiline_text((x+4,y+3),'\n'.join(lines),font=font,fill=color,spacing=2)
            d.rectangle((x,y,x+width-1,y+h-1),outline='#dddddd');x+=width
        y+=h
    im.save(path)
render(s,out/'manifest-before.png')
if not a.edit:print('Preview created');raise SystemExit(0)
dest=out/'Loan_Daily_Analysis_Template.xlsx'
with ZipFile(source) as original,ZipFile(dest,'w') as changed:
    for entry in original.infolist():
        data=original.read(entry.filename)
        if entry.filename=='xl/worksheets/sheet1.xml':
            text=data.decode()
            for address,value in changes.items():
                pattern=r'(<c r="'+address+r'"[^>]*><v>)(.*?)(</v></c>)'
                matches=re.findall(pattern,text);assert len(matches)==1
                assert matches[0][1]==escape(str(s[address].value))
                text=re.sub(pattern,lambda m:m[1]+escape(value)+m[3],text,count=1)
            data=text.encode()
        changed.writestr(entry,data)
v=load_workbook(dest);diff=[]
assert w.sheetnames==v.sheetnames
for before,after in zip(w,v):
    for row in before:
        for c in row:
            d=after[c.coordinate]
            if c.value!=d.value:diff.append((before.title,c.coordinate,c.value,d.value))
            assert c._style==d._style
assert {(x[0],x[1],x[3]) for x in diff}=={('Manifest',k,value) for k,value in changes.items()}
with ZipFile(source) as one,ZipFile(dest) as two:
    assert one.namelist()==two.namelist()
    changed_parts=[n for n in one.namelist() if one.read(n)!=two.read(n)]
assert changed_parts==['xl/worksheets/sheet1.xml']
render(v['Manifest'],out/'manifest-after.png')
result={'changed_cells':diff,'changed_zip_members':changed_parts,'all_other_workbook_parts_identical':True,
 'before_sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'after_sha256':hashlib.sha256(dest.read_bytes()).hexdigest(),
 'fallback_reason':'Bundled @oai/artifact-tool package absent at configured runtime and repository junction.',
 'runtime_template_rendering':'Not executed; static builtin syntax verified against official AppSheet documentation.'}
(out/'verification.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result))
