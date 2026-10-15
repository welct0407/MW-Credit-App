"""Read-only verification of canonical workbook scope and native properties."""
import copy
import json
from pathlib import Path
import openpyxl

root = Path('C:/GitHub/AppSheet-Loan-Project')
before = openpyxl.load_workbook('C:/Users/MWCredit/Documents/ChatGPT/AppSheet-Loan-Project/lumpsum-20260915/dictionary-before.xlsx')
after = openpyxl.load_workbook(root / 'Documents/AppSheet_Data_Dictionary.xlsx')
assert before.sheetnames == after.sheetnames
assert list(before.defined_names) == list(after.defined_names)
checked = 0
for name in before.sheetnames:
    a, b = before[name], after[name]
    assert a.freeze_panes == b.freeze_panes, (name, 'panes')
    assert a.auto_filter == b.auto_filter, (name, 'filter')
    assert list(a.tables) == list(b.tables), (name, 'tables')
    assert len(a._charts) == len(b._charts), (name, 'charts')
    assert str(a.merged_cells) == str(b.merged_cells), (name, 'merges')
    assert str(a.data_validations) == str(b.data_validations), (name, 'validation')
    assert str(a.protection) == str(b.protection), (name, 'protection')
    for row in a:
        for cell in row:
            target = b[cell.coordinate]
            allowed = name == 'App Settings' and a.cell(cell.row, 2).value == 'Sync' and cell.column in (4, 5)
            if not allowed:
                assert cell.value == target.value, (name, cell.coordinate, 'value')
                assert cell.data_type == target.data_type, (name, cell.coordinate, 'type')
            for attr in ('font', 'fill', 'border', 'alignment', 'number_format', 'protection'):
                assert copy.copy(getattr(cell, attr)) == copy.copy(getattr(target, attr)), (name, cell.coordinate, attr)
            checked += 1
result = dict(sheets=len(before.sheetnames),existingCellsChecked=checked,existingValuesAndStylesPreserved=True,
              intendedExceptions='App Settings sync Value/Evidence; appended development documentation rows',nativePropertiesPreserved=True)
(root / 'outputs/lumpsum-db-20260915/dictionary-native-verification.json').write_text(json.dumps(result, indent=2)+'\n', encoding='utf-8')
print(json.dumps(result))
