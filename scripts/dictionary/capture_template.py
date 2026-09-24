"""Extract a downloaded, metadata-only XLSX template and validate literal references.

Never renders templates, evaluates expressions or accesses business rows.
"""
import argparse
import datetime as dt
import json
import re
from pathlib import Path
from openpyxl import load_workbook
from common import digest, write_json

def capture(path, folder):
    path=Path(path);folder=Path(folder)
    workbook=load_workbook(path,read_only=True,data_only=False)
    sheets=[]
    for sheet in workbook:
        cells=[{'address':cell.coordinate,'type':cell.data_type,'value':cell.value}
               for row in sheet.iter_rows() for cell in row if cell.value is not None]
        sheets.append({'name':sheet.title,'cells':cells})
    workbook.close()
    result={'id':'1g4xHHsTVvRR5ek4u3DslZAb_UerH05E5',
      'url':'https://drive.google.com/file/d/1g4xHHsTVvRR5ek4u3DslZAb_UerH05E5/view',
      'source_modified_utc':'2026-08-31T02:45:22.635Z',
      'captured_utc':dt.datetime.now(dt.timezone.utc).isoformat(),
      'method':'Google Drive connector raw stored XLSX; openpyxl read-only cell extraction',
      'sha256':digest(path.read_bytes()),'sheets':sheets}
    findings=[];checks={}
    contexts={'Statistics':'Reporting Statistics','Daily Analytics 30D':'Reporting Daily Analytics',
      'Borrowers':'Reporting Borrowers','Loans':'Reporting Loans','Charges':'Reporting Charges','Repayments 30D':'Reporting Repayments'}
    for env in ['production','development']:
        source=json.loads((folder/(env+'-olap-observations.json')).read_text(encoding='utf8'))
        tables={x['name']:set(x['content']['columns']) for x in source['objects'] if x['kind']=='table'}
        for obj in source['objects']:
            if obj['kind']=='slice':tables[obj['name']]=set(obj['content']['definition']['Slice Columns'])
        checked=0
        for sheet in sheets:
            for cell in sheet['cells']:
                text=str(cell['value'])
                # Only direct output expressions in a known Start context. Complex
                # SELECT predicates remain exact captured text, not parser claims.
                if sheet['name'] in contexts:
                    for column in re.findall(r'<<\[([^\]]+)\]>>',text):
                        checked+=1
                        if column not in tables[contexts[sheet['name']]]:
                            findings.append({'target':env+'-olap','sheet':sheet['name'],'cell':cell['address'],'missing_column':column,'slice':contexts[sheet['name']]})
                for table,column in re.findall(r'(Reporting [A-Za-z ]+)\[([^\]]+)\]',text):
                    checked+=1
                    if table not in tables or column not in tables[table]:
                        findings.append({'target':env+'-olap','sheet':sheet['name'],'cell':cell['address'],'missing_reference':table+'['+column+']'})
        checks[env+'-olap']=checked
    write_json(folder/'export-template-observations.json',result)
    write_json(folder/'template-validation.json',{'checked_literal_references':checks,'findings':findings,
        'limitations':['No template rendering or complex expression execution. Workbook styling is outside data-dictionary scope.']})
    print(json.dumps({'sheets':len(sheets),'cells':sum(len(s['cells']) for s in sheets),'references':checks,'findings':findings}))

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('path');p.add_argument('folder');a=p.parse_args();capture(a.path,a.folder)
