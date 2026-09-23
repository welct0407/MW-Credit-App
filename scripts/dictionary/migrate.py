"""One-time, lossless reorganization of workbook records. No reconciliation."""
import argparse
import collections
import datetime
import json
import re
from pathlib import Path

import openpyxl
from common import REPO, ROOT, digest, equal, escape, make_index, write_json, write_yaml

SOURCE_ID = "excel-20260923"
OUTPUT = REPO / "outputs/dictionary-migration-20260923"


def slug(value):
    text = str(value) if value is not None else "unassigned"
    result = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-") or "unnamed"
    if result in {"con", "prn", "aux", "nul", *(f"com{i}" for i in range(1, 10)), *(f"lpt{i}" for i in range(1, 10))}:
        result = "object-" + result
    return result if len(result) <= 75 else result[:62].rstrip("-") + "-" + digest(text.encode())[:12]


def cell_value(cell, cached):
    if cell.data_type == "f":
        return {"excel_formula": cell.value, "cached_value": cached.value}
    if isinstance(cell.value, (datetime.datetime, datetime.date, datetime.time)):
        return {"excel_date_type": type(cell.value).__name__, "value": cell.value.isoformat()}
    if cell.data_type == "e":
        return {"excel_error": cell.value}
    return cell.value


def read_workbook(path):
    workbook = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    sheets = {}
    features = {"formulas": [], "comments": [], "hyperlinks": [], "hidden_rows": [], "hidden_columns": [], "merged_ranges": []}
    for sheet in workbook:
        headers = [c.value for c in sheet[1]]
        if any(not isinstance(h, str) for h in headers) or len(set(headers)) != len(headers):
            raise ValueError(f"Missing or duplicate headers in {sheet.title}")
        rows = []
        for row in sheet.iter_rows(min_row=2):
            values = [cell_value(c, cached[sheet.title][c.coordinate]) for c in row]
            if any(v is not None for v in values):
                rows.append((row[0].row, values))
            for c in row:
                if c.data_type == "f":
                    features["formulas"].append([sheet.title, c.coordinate, cell_value(c, cached[sheet.title][c.coordinate])])
                if c.comment:
                    features["comments"].append([sheet.title, c.coordinate, c.comment.text, c.comment.author])
                if c.hyperlink:
                    features["hyperlinks"].append([sheet.title, c.coordinate, c.hyperlink.target, c.hyperlink.location])
        features["hidden_rows"].extend([sheet.title, n] for n, d in sheet.row_dimensions.items() if d.hidden)
        features["hidden_columns"].extend([sheet.title, n] for n, d in sheet.column_dimensions.items() if d.hidden)
        features["merged_ranges"].extend([sheet.title, str(r)] for r in sheet.merged_cells.ranges)
        sheets[sheet.title] = {"headers": headers, "rows": rows, "sheet_state": sheet.sheet_state}
    return sheets, features


class Migration:
    def __init__(self):
        self.docs = {}
        self.locations = []
        self.unplaced = []

    def doc(self, file, system, component, kind, name, scope):
        identity = dict(system=system, component=component, kind=kind, name=name, scope=scope)
        if file not in self.docs:
            self.docs[file] = {"format": 1, "identity": identity,
                               "verification": {"status": "imported-unreconciled", "source": SOURCE_ID}, "content": {}}
        if not equal(self.docs[file]["identity"], identity):
            raise ValueError(f"Filename collision: {file}")
        return self.docs[file]["content"]

    def record(self, sheet, number, file, pointer, data, **location):
        self.locations.append({"sheet": sheet, "row": number, "file": file, "pointer": pointer,
                               "shape": "record", **location})
        return data

    def generic(self, sheet, number, record, group, file=None, system="workbook-evidence", component=SOURCE_ID, kind="sheet-records"):
        file = file or f"imports/{SOURCE_ID}/{slug(sheet)}/{slug(group)}.yaml"
        if system == "workbook-evidence":
            component = SOURCE_ID + "/" + slug(sheet)
        content = self.doc(file, system, component, kind, str(group) if group is not None else "(unassigned)", "as-recorded")
        content.setdefault("sheet", sheet)
        if content["sheet"] != sheet:
            raise ValueError("Generic group mixed sheets")
        key = f"row-{number:05d}"
        content.setdefault("records", {})[key] = self.record(sheet, number, file, "/content/records/" + key, record)

    def migrate_row(self, sheet, number, headers, values):
        record = dict(zip(headers, values))
        if sheet in ("Prod Tables", "Prod Columns", "Column Rules", "Prod Relationships"):
            app = record.get("App")
            table = record.get("Table", record.get("Child table"))
            if not app or not table:
                self.generic(sheet, number, record, "unplaced")
                self.unplaced.append([sheet, number])
                return
            file = f"systems/appsheet/{slug(app)}/tables/{slug(table)}.yaml"
            content = self.doc(file, "appsheet", slug(app), "table", table, "recorded-production")
            content.setdefault("source_app", app)
            content.setdefault("source_table", table)
            if sheet == "Prod Tables":
                if "binding" in content:
                    raise ValueError(f"Duplicate table binding: {app}/{table}")
                content["binding"] = self.record(sheet, number, file, "/content/binding", record)
            else:
                column = record.get("Column", record.get("Ref column"))
                if not isinstance(column, str):
                    raise ValueError("Column identity missing")
                col = content.setdefault("columns", {}).setdefault(column, {})
                base = "/content/columns/" + escape(column)
                if sheet == "Column Rules":
                    prop = record["Property"]
                    rules = col.setdefault("rules", {})
                    if not isinstance(prop, str) or prop in rules:
                        raise ValueError("Missing/duplicate rule identity")
                    rules[prop] = record["Definition"]
                    self.locations.append({"sheet": sheet, "row": number, "file": file,
                                           "pointer": base + "/rules/" + escape(prop), "shape": "rule",
                                           "column": column, "property": prop})
                else:
                    key = "inventory" if sheet == "Prod Columns" else "relationship"
                    if key in col:
                        raise ValueError("Duplicate column inventory/relationship")
                    col[key] = self.record(sheet, number, file, base + "/" + key, record)
        elif sheet in ("Prod Slices", "Slice Columns"):
            app, name = record["App"], record["Slice"]
            if not app or not name:
                self.generic(sheet, number, record, "unplaced")
                self.unplaced.append([sheet, number])
                return
            file = f"systems/appsheet/{slug(app)}/slices/{slug(name)}.yaml"
            content = self.doc(file, "appsheet", slug(app), "slice", name, "recorded-production")
            if sheet == "Prod Slices":
                if "definition" in content:
                    raise ValueError("Duplicate slice definition")
                content["definition"] = self.record(sheet, number, file, "/content/definition", record)
            else:
                members = content.setdefault("members", [])
                members.append(self.record(sheet, number, file, f"/content/members/{len(members)}", record))
        elif sheet in ("DB Columns", "DB Constraints"):
            name = record.get("Table / view", record.get("Table"))
            file = f"systems/postgresql/relations/{slug(name)}.yaml"
            content = self.doc(file, "postgresql", "unassigned-database", "relation", name, "as-recorded")
            section, key = ("columns", record["Column"]) if sheet == "DB Columns" else ("constraints", record["Constraint"])
            group = content.setdefault(section, {})
            if key in group:
                raise ValueError("Duplicate SQL column/constraint")
            group[key] = self.record(sheet, number, file, f"/content/{section}/" + escape(key), record)
        elif sheet == "DB Logic":
            group = record["Schema / table"]
            self.generic(sheet, number, record, group, f"systems/postgresql/recorded-logic/{slug(group)}.yaml",
                         "postgresql", "unassigned-database", "recorded-logic")
        elif sheet == "Metabase Models":
            name, field = record["Model"], record["Field / grain"]
            file = f"systems/metabase/models/{slug(name)}.yaml"
            content = self.doc(file, "metabase", "recorded-instance", "model", name, "as-recorded")
            fields = content.setdefault("fields", {})
            if field in fields:
                raise ValueError("Duplicate Metabase field")
            fields[field] = self.record(sheet, number, file, "/content/fields/" + escape(field), record)
        elif sheet == "App Settings":
            group = record["Environment / app"]
            self.generic(sheet, number, record, group, f"systems/appsheet/recorded-settings/{slug(group)}.yaml",
                         "appsheet", "recorded-labels", "settings")
        elif sheet == "Dev Differences":
            self.generic(sheet, number, record, record["App"])
        elif sheet == "R008 Production":
            self.generic(sheet, number, record, record["App"])
        elif sheet == "Guide":
            self.generic(sheet, number, record, "guide")
        else:
            raise ValueError(f"Unmapped worksheet: {sheet}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=REPO / "Documents/AppSheet_Data_Dictionary.xlsx")
    parser.add_argument("--root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    if args.root.exists() or args.output.exists():
        raise ValueError("Migration target/evidence directory already exists; refusing overwrite")
    source_hash = digest(args.source.read_bytes())
    sheets, features = read_workbook(args.source)
    migration = Migration()
    for name, sheet in sheets.items():
        for number, values in sheet["rows"]:
            migration.migrate_row(name, number, sheet["headers"], values)
    if digest(args.source.read_bytes()) != source_hash:
        raise ValueError("Source workbook changed during extraction")
    args.root.mkdir(parents=True)
    args.output.mkdir(parents=True)
    for file, doc in migration.docs.items():
        write_yaml(args.root / file, doc)
    registry = {"format": 1, "migration": {"id": SOURCE_ID, "workbook": "../AppSheet_Data_Dictionary.xlsx", "sha256": source_hash},
                "systems": {
                    "appsheet": {"status": "imported-unreconciled", "kinds": ["table", "slice", "settings", "action", "view", "automation"]},
                    "postgresql": {"status": "imported-unreconciled", "kinds": ["relation", "recorded-logic", "routine"]},
                    "metabase": {"status": "imported-unreconciled", "kinds": ["model", "question", "dashboard", "snippet"]},
                    "grafana": {"status": "planned", "kinds": ["dashboard", "datasource", "alert-rule", "recording-rule"]},
                    "workbook-evidence": {"status": "imported-unreconciled", "kinds": ["sheet-records"]},
                }}
    write_yaml(args.root / "catalog.yaml", registry)
    issues = [
        {"id": "UNRECONCILED", "status": "deferred", "detail": "Every imported definition retains the workbook's recorded scope and value. No live inspection, precedence decision or reconciliation was performed; owner will reconcile later."},
        {"id": "SUPPLEMENTS", "status": "deferred", "detail": "R008 Production contains unique definitions and accepted targets. Imported separately; they are not applied as overrides to base definitions.", "path": f"imports/{SOURCE_ID}/r008-production"},
        {"id": "DEV-HISTORY", "status": "deferred", "detail": "Dev Differences mixes dated, closed and potentially active records. Labels are preserved verbatim. These records must not be automatically applied as current DEV overrides.", "path": f"imports/{SOURCE_ID}/dev-differences"},
        {"id": "UNPLACED", "status": "preserved", "detail": "Seven nonempty rows have incomplete object identities. They remain text records rather than invented objects.", "source_rows": migration.unplaced},
        {"id": "LABELS", "status": "deferred", "detail": "Different environment labels, removed-app labels and property spellings remain separate. Database/Metabase instance identity is unassigned where the workbook does not establish it."},
        {"id": "EXCEL-FORMULA", "status": "preserved", "detail": "App Settings D121 is an Excel formula. Both exact formula text and its stored cached value are preserved; no expression evaluation or reinterpretation was performed."},
        {"id": "COVERAGE", "status": "deferred", "detail": "This is a complete migration of the workbook's records, not a complete inventory of the ecosystem. Grafana is a planned namespace with no imported definitions. AppSheet action/view/automation details remain in their source supplements until later reconciliation."},
    ]
    write_yaml(args.root / "unresolved.yaml", {"format": 1, "items": issues})
    write_json(args.root / "index.json", make_index(args.root))
    manifest = {"format": 1, "source": "Documents/AppSheet_Data_Dictionary.xlsx", "sha256": source_hash,
                "source_id": SOURCE_ID, "sheets": {name: {"headers": s["headers"], "sheet_state": s["sheet_state"],
                "rows": len(s["rows"]), "row_numbers": [n for n, _ in s["rows"]]} for name, s in sheets.items()},
                "features": features, "documents": len(migration.docs), "records": len(migration.locations),
                "files": {file: digest((args.root / file).read_bytes()) for file in sorted(migration.docs)}}
    write_json(args.output / "source-manifest.json", manifest)
    # Compact, one source row per line: searchable migration evidence without duplicated values.
    with (args.output / "row-map.jsonl").open("w", encoding="utf-8", newline="\n") as out:
        for row in migration.locations:
            out.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
    print(json.dumps({"documents": len(migration.docs), "records": len(migration.locations), "source_sha256": source_hash,
                      "formula_cells": len(features["formulas"]), "unplaced_rows": migration.unplaced}))


if __name__ == "__main__":
    main()
