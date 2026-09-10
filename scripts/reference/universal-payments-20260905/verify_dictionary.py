import openpyxl
from pathlib import Path
root=Path(__file__).parent
a=openpyxl.load_workbook(root/'AppSheet_Data_Dictionary.before.xlsx')
b=openpyxl.load_workbook(root/'AppSheet_Data_Dictionary.xlsx')
issues=[]
for name in a.sheetnames:
 x,y=a[name],b[name]
 for row in x:
  for c in row:
   d=y[c.coordinate]
   if c.value!=d.value: issues.append((name,c.coordinate,'value changed'))
   if c.number_format!=d.number_format: issues.append((name,c.coordinate,'number format changed'))
   if c.has_style and (c.font!=d.font or c.fill!=d.fill or c.alignment!=d.alignment or c.border!=d.border):
    # Component comparisons must compare actual style objects rather than proxy identity.
    from copy import copy
    if any(copy(getattr(c,p))!=copy(getattr(d,p)) for p in ['font','fill','alignment','border']): issues.append((name,c.coordinate,'style changed'))
 if x.freeze_panes!=y.freeze_panes: issues.append((name,'freeze panes changed'))
 if x.sheet_state!=y.sheet_state: issues.append((name,'visibility changed'))
 if not set(str(z) for z in x.merged_cells.ranges)<=set(str(z) for z in y.merged_cells.ranges): issues.append((name,'existing merge changed'))
 if len(x.data_validations.dataValidation)!=len(y.data_validations.dataValidation): issues.append((name,'validation changed'))
 if len(x._charts)!=len(y._charts): issues.append((name,'chart changed'))
 if [(t.name,t.ref) for t in x.tables.values()]!=[(t.name,t.ref) for t in y.tables.values()]: issues.append((name,'table changed'))
print({'original_sheets':len(a.sheetnames),'new_sheets':set(b.sheetnames)-set(a.sheetnames),'preservation_issues':len(issues),'examples':issues[:20]})
assert not issues
defs=[tuple(c.value for c in row) for row in b['Payments Clone'].iter_rows(min_row=5,max_row=58,max_col=8)]
assert len(defs)==54 and all(r[4]=='OLTP test clone only' for r in defs)
refs={(r[0],r[2]) for r in defs if r[3]=='Physical / Ref'}
expected_refs={('Payments','Ref Borrower'),('Payments','Ref Target Charge'),('Payment Allocations','Ref Payment'),('Payment Allocations','Ref Charge'),('Repayments','Ref Payment'),('Repayments','Ref Payment Allocation'),('Repayments','Ref Charges')}
assert refs==expected_refs, refs
assert {r[0]:sum(x[0]==r[0] for x in defs) for r in defs}=={'Payments':23,'Payment Allocations':19,'Repayments':8,'Borrowers':1,'Charges':3}
assert len([c for row in b['Payments Clone'] for c in row if c.data_type=='f'])==0
print('Existing workbook content and checked presentation structures preserved; 54 definitions, six new and one existing Reference verified; expressions stored as text.')

import json
expected=json.loads((root/'clone_slices_checkpoint065.json').read_text(encoding='utf-8'))
actual=[tuple(c.value for c in row) for row in b['Payments Clone Slices'].iter_rows(min_row=5,max_row=49,max_col=7)]
wanted=[(x['name'],x['source'],i+1,c,x['filter'],x['updates'],x['actions']) for x in expected for i,c in enumerate(x['columns'])]
assert actual==wanted
assert len(actual)==45
print('Both clone slices verified: Visible Borrowers33; Todays Collection Board12; production slice inventory preserved.')
