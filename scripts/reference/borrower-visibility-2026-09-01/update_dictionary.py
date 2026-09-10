from __future__ import annotations

import argparse
import shutil
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(r"D:\GitHub\AppSheet Project")
WORKBOOK = ROOT / "Documents" / "AppSheet_Data_Dictionary.xlsx"
OUTPUT_DIR = ROOT / "outputs" / "borrower-visibility-2026-09-01"
OUTPUT_WORKBOOK = OUTPUT_DIR / "AppSheet_Data_Dictionary.xlsx"


def copy_row_style(ws, source_row: int, target_row: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, ws.max_column + 1):
        source = ws.cell(source_row, col)
        target = ws.cell(target_row, col)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def find_row(ws, column: int, value: str, start: int = 1) -> int:
    for row in range(start, ws.max_row + 1):
        if ws.cell(row, column).value == value:
            return row
    raise ValueError(f"Could not find {value!r} in {ws.title} column {column}")


def set_readme_value(ws, control: str, value) -> None:
    row = find_row(ws, 1, control)
    ws.cell(row, 2).value = value


def update_workbook() -> None:
    wb = load_workbook(WORKBOOK)

    readme = wb["README"]
    set_readme_value(
        readme,
        "Reconciliation status",
        "Reconciled with production OLTP v2.200099 on 2026-09-01 for the Borrowers visibility control. "
        "Money database now contains Borrowers[Hidden Flag]; production OLTP loads it and filters only the primary Borrowers view through Visible Borrowers. "
        "OLAP remains at its prior structural baseline and has not regenerated this physical field.",
    )
    set_readme_value(
        readme,
        "Live baselines",
        "OLTP v2.200099 — Loan Manager ผู้จัดการสินเชื่อ — 93844270-6aa5-420b-a960-bcdf08abd357\n"
        "OLAP v1.000151 — Loan Business Insights — e74dd368-654f-44d7-bccf-83cf361d1cdb",
    )
    set_readme_value(readme, "Slices", 13)

    tables = wb["Tables"]
    borrower_table_row = find_row(tables, 1, "Borrowers", start=5)
    tables.cell(borrower_table_row, 4).value = "7"
    tables.cell(borrower_table_row, 7).value = (
        "Borrower master records and borrower-level operational or portfolio indicators; "
        "production OLTP also stores a view-only Hidden Flag."
    )

    columns = wb["Columns"]
    creation_row = next(
        row
        for row in range(5, columns.max_row + 1)
        if columns.cell(row, 1).value == "Borrowers"
        and columns.cell(row, 3).value == "Creation Date"
    )
    insert_row = creation_row + 1
    columns.insert_rows(insert_row, 1)
    copy_row_style(columns, creation_row, insert_row)

    row = insert_row + 1
    while row <= columns.max_row and columns.cell(row, 1).value == "Borrowers":
        order = columns.cell(row, 2).value
        if isinstance(order, (int, float)) and order >= 7:
            columns.cell(row, 2).value = order + 1
        row += 1

    hidden_column = [
        "Borrowers",
        7,
        "Hidden Flag",
        "Physical / Yes/No",
        "OLTP only",
        "OLTP: Show and editable; initial FALSE; searchable; OLAP: Not loaded",
        "View-only visibility control. TRUE hides the borrower from the production OLTP Borrowers list without excluding related loans, collections, or calculations.",
        'Stored source column; initial value: FALSE; display name: "ซ่อนจากรายชื่อผู้กู้ / Hide from Borrowers List"',
    ]
    for col, value in enumerate(hidden_column, start=1):
        columns.cell(insert_row, col).value = value
    columns.tables["DataModelColumns"].ref = f"A4:H{columns.max_row}"

    slices = wb["Slices"]
    today_slice_row = find_row(slices, 2, "Todays Collection Board", start=5)
    visible_slice_row = today_slice_row + 1
    slices.insert_rows(visible_slice_row, 1)
    copy_row_style(slices, today_slice_row, visible_slice_row)
    visible_slice = [
        "OLTP",
        "Visible Borrowers",
        "Borrowers",
        "NOT([Hidden Flag])",
        "Adds, updates, deletes",
        31,
        "Editable production borrower-list slice that excludes only rows explicitly marked hidden.",
        "Primary OLTP borrower-list slice. Excludes rows where [Hidden Flag] is TRUE; base-table relationships, loans, collections, and calculations remain unchanged.",
    ]
    for col, value in enumerate(visible_slice, start=1):
        slices.cell(visible_slice_row, col).value = value
    slices.tables["DataModelSlices"].ref = f"A4:H{slices.max_row}"

    slice_columns = wb["Slice Columns"]
    today_rows = [
        row
        for row in range(5, slice_columns.max_row + 1)
        if slice_columns.cell(row, 1).value == "OLTP"
        and slice_columns.cell(row, 2).value == "Todays Collection Board"
    ]
    visible_start = max(today_rows) + 1
    visible_columns = [
        "_RowNumber",
        "Row ID",
        "Borrower Name",
        "Telephone",
        "Description",
        "Creation Date",
        "Hidden Flag",
        "Related Loans",
        "Total Interest Earned",
        "Total Amount Loaned",
        "Total Number of Loans",
        "Total Outstanding Principal",
        "Header - Active Loans",
        "Borrower Name TE",
        "Return on Active Principal",
        "Active Loan Interest Earned",
        "Active Daily Interest",
        "Projected Interest Coverage Date",
        "Closed Loan Interest Earned",
        "Closed Loan Principle",
        "Has Closed Loan",
        "Return on Closed Loans",
        "Header - Closed Loans",
        "Has Active Loan",
        "Daily Collection Status",
        "Todays Collection Charges",
        "Todays Amount Due",
        "Todays Amount Collected",
        "Todays Amount Remaining",
        "Daily Collection Status - Order",
        "Borrower Name - Collection",
    ]
    slice_columns.insert_rows(visible_start, len(visible_columns))
    for offset, name in enumerate(visible_columns):
        target_row = visible_start + offset
        copy_row_style(slice_columns, visible_start - 1, target_row)
        values = ["OLTP", "Visible Borrowers", offset + 1, name]
        for col, value in enumerate(values, start=1):
            slice_columns.cell(target_row, col).value = value
    slice_columns.tables["DataModelSliceColumns"].ref = f"A4:D{slice_columns.max_row}"

    variances = wb["App Variances"]
    borrower_variance_rows = [
        row
        for row in range(5, variances.max_row + 1)
        if variances.cell(row, 2).value == "Borrowers"
    ]
    variance_row = max(borrower_variance_rows) + 1
    variances.insert_rows(variance_row, 1)
    copy_row_style(variances, variance_row - 1, variance_row)
    variance_values = [
        "Column",
        "Borrowers",
        "Hidden Flag",
        "Yes",
        "No",
        "Physical Yes/No field exists in Money database and production OLTP. OLAP schema was not regenerated; the field is used only by the OLTP Visible Borrowers slice.",
    ]
    for col, value in enumerate(variance_values, start=1):
        variances.cell(variance_row, col).value = value
    variances.tables["DataModelAppVariances"].ref = f"A4:F{variances.max_row}"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(WORKBOOK)
    shutil.copy2(WORKBOOK, OUTPUT_WORKBOOK)


def color_from_cell(cell, default: str = "#FFFFFF") -> str:
    fill = cell.fill
    if fill and fill.fill_type and fill.fgColor.type == "rgb" and fill.fgColor.rgb:
        rgb = fill.fgColor.rgb[-6:]
        return f"#{rgb}"
    return default


def load_font(size: int, bold: bool = False):
    candidates = [
        Path(r"C:\Windows\Fonts\leelawui.ttf"),
        Path(r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def abbreviate(value, max_chars: int = 180) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r", " ").replace("\n", " ⏎ ")
    if len(text) > max_chars:
        return text[: max_chars - 1] + "…"
    return text


def wrap_text(draw, text: str, font, width: int, max_lines: int = 3) -> list[str]:
    if not text:
        return []
    words = text.split(" ")
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else current + " " + word
        if draw.textlength(candidate, font=font) <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
            if len(lines) >= max_lines:
                break
    if current and len(lines) < max_lines:
        lines.append(current)
    if len(lines) == max_lines and words:
        lines[-1] = lines[-1][: max(1, len(lines[-1]) - 1)] + "…"
    return lines


def render_sheet(ws, out_path: Path) -> None:
    row_limits = {
        "README": 26,
        "Tables": 15,
        "Columns": 22,
        "Relationships": 12,
        "Slices": 17,
        "Slice Columns": 52,
        "App Variances": 45,
        "Change Log": 6,
    }
    max_row = min(ws.max_row, row_limits.get(ws.title, 25))
    max_col = ws.max_column
    natural_widths = []
    for col in range(1, max_col + 1):
        configured = ws.column_dimensions[get_column_letter(col)].width or 12
        natural_widths.append(max(90, min(300, int(configured * 7.2))))
    target_width = 1500
    scale = min(1.0, (target_width - 30) / sum(natural_widths))
    widths = [max(70, int(width * scale)) for width in natural_widths]

    row_heights = []
    for row in range(1, max_row + 1):
        configured = ws.row_dimensions[row].height or 20
        row_heights.append(max(24, min(62, int(configured * 1.2))))
    width = sum(widths) + 30
    height = sum(row_heights) + 55
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(22, bold=True)
    normal_font = load_font(12)
    bold_font = load_font(12, bold=True)
    draw.text((14, 10), ws.title, fill="#16324F", font=title_font)

    y = 48
    for row in range(1, max_row + 1):
        x = 14
        row_height = row_heights[row - 1]
        for col in range(1, max_col + 1):
            cell = ws.cell(row, col)
            cell_width = widths[col - 1]
            fill = color_from_cell(cell, "#FFFFFF")
            if row == 4:
                fill = "#1F4E78"
            draw.rectangle((x, y, x + cell_width, y + row_height), fill=fill, outline="#B7C4CE")
            text = abbreviate(cell.value)
            font = bold_font if row in (1, 4) or cell.font.bold else normal_font
            font_color = "#FFFFFF" if row == 4 else "#17212B"
            lines = wrap_text(draw, text, font, cell_width - 10, max_lines=3)
            ty = y + 5
            for line in lines:
                draw.text((x + 5, ty), line, fill=font_color, font=font)
                ty += 15
            x += cell_width
        y += row_height
    image.save(out_path)


def render_workbook(phase: str) -> None:
    wb = load_workbook(WORKBOOK, data_only=False)
    render_dir = OUTPUT_DIR / f"render-{phase}"
    render_dir.mkdir(parents=True, exist_ok=True)
    previews = []
    for index, ws in enumerate(wb.worksheets, start=1):
        path = render_dir / f"{index:02d}-{ws.title.replace(' ', '_')}.png"
        render_sheet(ws, path)
        previews.append((ws.title, Image.open(path).convert("RGB")))

    thumb_width, thumb_height = 1000, 690
    contact = Image.new("RGB", (thumb_width * 2, thumb_height * 4), "#E8EEF3")
    draw = ImageDraw.Draw(contact)
    label_font = load_font(18, bold=True)
    for index, (title, preview) in enumerate(previews):
        preview.thumbnail((thumb_width - 30, thumb_height - 45))
        col = index % 2
        row = index // 2
        x = col * thumb_width + 15
        y = row * thumb_height + 30
        draw.text((x, row * thumb_height + 5), title, fill="#16324F", font=label_font)
        contact.paste(preview, (x, y))
    contact.save(render_dir / "00-contact-sheet.png")


def validate() -> None:
    wb = load_workbook(WORKBOOK, data_only=False)
    columns = wb["Columns"]
    slices = wb["Slices"]
    slice_columns = wb["Slice Columns"]
    variances = wb["App Variances"]

    hidden_rows = [
        row
        for row in columns.iter_rows(min_row=5, values_only=True)
        if row[0] == "Borrowers" and row[2] == "Hidden Flag"
    ]
    assert len(hidden_rows) == 1
    assert hidden_rows[0][1] == 7
    assert hidden_rows[0][3] == "Physical / Yes/No"
    assert hidden_rows[0][4] == "OLTP only"

    visible_slices = [
        row
        for row in slices.iter_rows(min_row=5, values_only=True)
        if row[0] == "OLTP" and row[1] == "Visible Borrowers"
    ]
    assert len(visible_slices) == 1
    assert visible_slices[0][3] == "NOT([Hidden Flag])"
    assert visible_slices[0][5] == 31

    visible_membership = [
        row
        for row in slice_columns.iter_rows(min_row=5, values_only=True)
        if row[0] == "OLTP" and row[1] == "Visible Borrowers"
    ]
    assert len(visible_membership) == 31
    assert [row[2] for row in visible_membership] == list(range(1, 32))
    assert visible_membership[6][3] == "Hidden Flag"

    variance_rows = [
        row
        for row in variances.iter_rows(min_row=5, values_only=True)
        if row[1] == "Borrowers" and row[2] == "Hidden Flag"
    ]
    assert len(variance_rows) == 1
    assert variance_rows[0][3:5] == ("Yes", "No")

    assert columns.tables["DataModelColumns"].ref == f"A4:H{columns.max_row}"
    assert slices.tables["DataModelSlices"].ref == f"A4:H{slices.max_row}"
    assert slice_columns.tables["DataModelSliceColumns"].ref == f"A4:D{slice_columns.max_row}"
    assert variances.tables["DataModelAppVariances"].ref == f"A4:F{variances.max_row}"

    error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?")
    found_errors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and any(token in cell.value for token in error_tokens):
                    found_errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
    assert not found_errors, found_errors

    print(
        {
            "workbook": str(WORKBOOK),
            "output_copy": str(OUTPUT_WORKBOOK),
            "sheet_count": len(wb.sheetnames),
            "tables": wb["Tables"].max_row - 4,
            "columns": columns.max_row - 4,
            "slices": slices.max_row - 4,
            "slice_columns": slice_columns.max_row - 4,
            "app_variances": variances.max_row - 4,
            "formula_errors": found_errors,
        }
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["render-before", "update", "render-after", "validate"])
    args = parser.parse_args()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if args.mode == "render-before":
        render_workbook("before")
    elif args.mode == "update":
        update_workbook()
    elif args.mode == "render-after":
        render_workbook("after")
    else:
        validate()


if __name__ == "__main__":
    main()
