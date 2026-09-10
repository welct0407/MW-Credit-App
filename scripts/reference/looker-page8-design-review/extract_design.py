import sys

from docx import Document

sys.stdout.reconfigure(encoding="utf-8")

path = r"C:\Users\zhcyi\OneDrive\Documents\Looker_Studio_Page_8_Profit_Forecast_Design.docx"
doc = Document(path)

print("PARAGRAPHS")
for i, paragraph in enumerate(doc.paragraphs, start=1):
    text = paragraph.text.strip()
    if text:
        print(f"P{i}: {text}")

print("TABLES")
for ti, table in enumerate(doc.tables, start=1):
    print(f"TABLE {ti}: {len(table.rows)}x{len(table.columns)}")
    for ri, row in enumerate(table.rows, start=1):
        values = [cell.text.replace("\n", " | ").strip() for cell in row.cells]
        print(f"R{ri}: " + " || ".join(values))

print("SECTIONS")
for si, section in enumerate(doc.sections, start=1):
    print(
        f"S{si}: width={section.page_width} height={section.page_height} "
        f"margins={section.top_margin},{section.right_margin},{section.bottom_margin},{section.left_margin}"
    )
