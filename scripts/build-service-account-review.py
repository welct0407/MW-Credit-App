"""Temporary owner-input XLSX. Artifact Tool is absent from the bundled runtime.
Uses bundled XlsxWriter; never overwrite a workbook the owner may have filled.
"""
from pathlib import Path
import json
import xlsxwriter
import openpyxl
from PIL import Image, ImageDraw, ImageFont

root = Path(__file__).resolve().parents[1]
out = root / 'outputs/service-account-review-20260924'
out.mkdir(parents=True, exist_ok=True)
target = out / 'Service_Account_Review.xlsx'
if target.exists():
    raise SystemExit('Workbook already exists; preserve owner input.')
source = (root/'Documents/Service_Account_Browser_Mapping.md').read_text(encoding='utf-8')
section = source.split('## Service register')[1].split('## Owner input')[0]
services = [tuple(c.strip() for c in line.split('|')[1:3]) for line in section.splitlines()
            if line.startswith('| ') and not line.startswith(('| Service', '| ---'))]
assert len(services) == 17
first, last = 8, 29
workbook = xlsxwriter.Workbook(target)
workbook.set_calc_mode('auto')
sheet = workbook.add_worksheet('Service accounts')
sheet.hide_gridlines(2)
sheet.freeze_panes(7, 0)
widths = [218, 292, 278, 210, 242]
base = {'font_name':'Arial','font_size':11,'font_color':'#1F2937','valign':'vcenter','text_wrap':True}
body = workbook.add_format(base)
account = workbook.add_format({**base,'bg_color':'#FFF2CC'})
notes = workbook.add_format({**base,'bg_color':'#FFF9E8'})
browser = workbook.add_format({**base,'bg_color':'#F0F4F8'})
header = workbook.add_format({**base,'bold':True,'font_color':'#FFFFFF','bg_color':'#17324D','align':'center','right':1,'right_color':'#FFFFFF'})
title = workbook.add_format({'font_name':'Arial','font_size':16,'bold':True,'font_color':'#17324D'})
hint = workbook.add_format({'font_name':'Arial','font_size':10,'font_color':'#475569'})
for col,width in enumerate(widths):
    sheet.set_column_pixels(col,col,width,body)
sheet.set_row_pixels(0,8)
sheet.set_row_pixels(1,32)
sheet.write(1,0,'Service accounts',title)
instructions = [
 'Fill yellow Account cells with the actual login. Add notes or split rows for different roles. Browser is calculated.',
 'Chrome work = welct0407@mw-credit.com   |   Chrome Gmail = welct0407@gmail.com   |   Other accounts = Edge',
 'Services: Environment Inventory, Solution Design, Resource Monitoring and Git remote. Account assignments await your input.'
]
for row,text in enumerate(instructions,2):
    sheet.set_row_pixels(row,23)
    sheet.write(row,0,text,hint)
sheet.set_row_pixels(5,9)
sheet.set_row_pixels(6,32)
sheet.write_row(6,0,['Service','Scope / environment','Account (fill in)','Browser (calculated)','Notes (optional)'],header)
for i in range(22):
    row,excel_row = i+7,i+8
    sheet.set_row_pixels(row,36)
    values = services[i] if i<len(services) else ('','')
    sheet.write(row,0,values[0],body)
    sheet.write(row,1,values[1],body)
    sheet.write_blank(row,2,None,account)
    formula = f'=IF(TRIM(C{excel_row})="","Pending",IF(LOWER(TRIM(C{excel_row}))="welct0407@mw-credit.com","Chrome work",IF(LOWER(TRIM(C{excel_row}))="welct0407@gmail.com","Chrome Gmail","Edge")))'
    sheet.write_formula(row,3,formula,browser,'Pending')
    sheet.write_blank(row,4,None,notes)
sheet.autofilter(6,0,last-1,4)
sheet.print_area(0,0,last-1,4)
sheet.set_landscape()
sheet.fit_to_pages(1,1)
workbook.close()

# Saved-file verification, including every routing reference and empty owner field.
formulas = openpyxl.load_workbook(target,data_only=False)
cached = openpyxl.load_workbook(target,data_only=True)
s,values = formulas.active,cached.active
assert formulas.sheetnames == ['Service accounts']
assert s.freeze_panes == 'A8'
for row in range(first,last+1):
    f=s.cell(row,4).value
    assert f.count(f'C{row}')==3
    assert 'welct0407@mw-credit.com' in f and 'welct0407@gmail.com' in f
    assert f.endswith(',"Edge")))')
    assert s.cell(row,3).value is None and values.cell(row,4).value=='Pending'
assert [(s.cell(i+8,1).value,s.cell(i+8,2).value) for i in range(17)]==services

# Render the saved cell values/styles. Native Excel evaluates routing on owner input.
heights=[8,32,23,23,23,9,32]+[36]*22
image=Image.new('RGB',(sum(widths)+1,sum(heights)+1),'white')
draw=ImageDraw.Draw(image)
def font(size=11,bold=False):
    return ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if bold else 'arial.ttf'),round(size*96/72))
def wrapped(text,width,face):
    result=[]
    for line in text.splitlines() or ['']:
        current=''
        for word in line.split():
            candidate=(current+' '+word).strip()
            if current and draw.textlength(candidate,font=face)>width:
                result.append(current); current=word
            else: current=candidate
        result.append(current)
    return result
y=0
for row,height in enumerate(heights,1):
    x=0
    for col,width in enumerate(widths,1):
        cell=s.cell(row,col); text=str(values.cell(row,col).value or '')
        if cell.fill.fgColor.type=='rgb' and cell.fill.patternType=='solid':
            draw.rectangle((x,y,x+width,y+height),fill='#'+cell.fill.fgColor.rgb[-6:])
        if text:
            face=font(cell.font.sz or 11,bool(cell.font.b))
            color='#'+cell.font.color.rgb[-6:] if cell.font.color and cell.font.color.type=='rgb' else '#1F2937'
            lines=wrapped(text,width-14,face) if row>=7 else [text]
            lineheight=face.size+2
            for j,line in enumerate(lines):
                dx=x+7 if row!=7 else x+(width-draw.textlength(line,font=face))/2
                draw.text((dx,y+(height-len(lines)*lineheight)/2+j*lineheight),line,font=face,fill=color)
        x+=width
    y+=height
image.save(out/'preview.png')
(out/'verification.json').write_text(json.dumps({
 'services':17,'extraRows':5,'accounts':'All blank for owner input',
 'verification':'Saved XLSX values, all formula references, routing constants, cached blank states, freeze pane and visual preview',
 'authoring':'Bundled XlsxWriter fallback: @oai/artifact-tool absent from bundled node_modules',
 'recalculation':'Automatic Excel calculation enabled; no native Excel/LibreOffice calculation engine available here'
},indent=2)+'\n',encoding='utf-8')
print(json.dumps({'output':str(target),'services':17,'blankRows':5,'savedFileChecks':'passed'}))
