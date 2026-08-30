import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

import build_appsheet_data_model as base


ROOT = Path(r"D:\GitHub\AppSheet Project")
SOURCE_DOCX = ROOT / "Documents" / "AppSheet_Data_Model_Reference.docx"
INVENTORY_JSON = ROOT / ".tmp" / "model_inventory.json"
CONVERSION_JSON = ROOT / ".tmp" / "conversion_data.json"
OUTPUT_DOCX = ROOT / "Documents" / "AppSheet_Data_Model_Overview.docx"

TABLE_NAMES = [
    "Borrowers",
    "Calendar",
    "Cash Pool Contributions",
    "Charges",
    "Daily Analytics",
    "Loan Assessment",
    "Loans",
    "Partners",
    "Repayments",
    "Settlements",
    "Statistics",
]


def matrix(table):
    return [[cell.text.strip() for cell in row.cells] for row in table.rows]


def parse_label(text):
    if ":" not in text:
        return "", text.strip()
    label, value = text.split(":", 1)
    return label.strip(), value.strip()


def make_conversion_data(source_doc, inventory):
    tables = source_doc.tables
    detailed = []
    for table_name, table in zip(TABLE_NAMES, tables[8:19]):
        rows = matrix(table)
        headers = rows[0]
        for order, values in enumerate(rows[1:], start=1):
            record = {headers[i]: values[i] for i in range(len(headers))}
            record["Table"] = table_name
            record["Column Order"] = order
            detailed.append(record)

    slice_columns = []
    slices = []
    for app_key, app_label in (("oltp", "OLTP"), ("olap", "OLAP")):
        for slice_def in inventory["apps"][app_key]["slices"]:
            mode = base.update_mode(slice_def)
            slices.append({
                "App": app_label,
                "Slice": slice_def["name"],
                "Source": slice_def["source"],
                "Row Filter": slice_def.get("filter") or "TRUE",
                "Update Mode": mode,
                "Column Count": len(slice_def.get("columns", [])),
                "Purpose": base.SLICE_PURPOSES.get(slice_def["name"], ""),
                "Comment": slice_def.get("comment", ""),
            })
            for order, column in enumerate(slice_def.get("columns", []), start=1):
                slice_columns.append({
                    "App": app_label,
                    "Slice": slice_def["name"],
                    "Column Order": order,
                    "Column": column,
                })

    app_variances = []
    table_rows = matrix(tables[4])[1:]
    for row in table_rows:
        loaded_in = row[2]
        if loaded_in != "OLTP + OLAP":
            app_variances.append({
                "Object Type": "Table",
                "Table": row[0],
                "Object": row[0],
                "OLTP Presence": "Yes" if "OLTP" in loaded_in else "No",
                "OLAP Presence": "Yes" if "OLAP" in loaded_in else "No",
                "Notes": row[6],
            })
    for row in detailed:
        availability = row["Availability"]
        if availability != "OLTP + OLAP":
            app_variances.append({
                "Object Type": "Column",
                "Table": row["Table"],
                "Object": row["Column"],
                "OLTP Presence": "Yes" if "OLTP" in availability else "No",
                "OLAP Presence": "Yes" if "OLAP" in availability else "No",
                "Notes": row["Purpose"],
            })

    paragraphs = source_doc.paragraphs
    return {
        "sourceFile": SOURCE_DOCX.name,
        "sourceBaselineDate": inventory.get("generatedAt", "2026-08-29"),
        "conversionDate": "2026-08-30",
        "conversionStatus": "Converted from the Word baseline; not reconciled with the live AppSheet configuration.",
        "documentControl": matrix(tables[0]),
        "architecture": matrix(tables[3]),
        "tables": matrix(tables[4]),
        "relationships": matrix(tables[5]),
        "slices": slices,
        "columns": detailed,
        "sliceColumns": slice_columns,
        "appVariances": app_variances,
        "implications": matrix(tables[19]),
        "sourceConfidence": [paragraphs[i].text for i in range(9, 13)],
        "updateProcedure": [paragraphs[i].text for i in range(14, 19)],
        "monthlyForecastAssumptions": [paragraphs[i].text for i in range(139, 147)],
        "manualChecklist": [paragraphs[i].text for i in range(150, 157)],
    }


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def copy_matrix_table(doc, source_matrix, widths, font_size=8.2):
    return base.make_table(doc, source_matrix[0], widths, source_matrix[1:], font_size=font_size)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.12
    r = p.add_run(text)
    base.set_font(r, size=9.3, color=base.INK)
    return p


def add_numbered_steps(doc, rows):
    for label, text in rows:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(label + ": ")
        base.set_font(r, size=9.5, bold=True, color=base.NAVY)
        r = p.add_run(text)
        base.set_font(r, size=9.5, color=base.INK)


def build_overview(data):
    doc = Document()
    doc.core_properties.title = "AppSheet Data Model Overview"
    doc.core_properties.subject = "Descriptive companion to the AppSheet technical data dictionary workbook"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "AppSheet, data model, OLTP, OLAP, governance, technical dictionary"
    base.configure_styles(doc)
    for section in doc.sections:
        base.configure_section(section)

    base.add_para(doc, "GOVERNED MODEL REFERENCE", size=10, bold=True, color=base.BLUE, after=22, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("AppSheet Data Model Overview")
    base.set_font(r, size=28, bold=True, color=base.NAVY)
    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Money database  |  Loan Manager ผู้จัดการสินเชื่อ  |  Loan Business Insights")
    base.set_font(r, size=13, color=base.MUTED)
    base.add_para(doc, "Descriptive companion to AppSheet_Data_Dictionary.xlsx", size=11, italic=True, color=base.MUTED, after=24, align=WD_ALIGN_PARAGRAPH.CENTER)
    cover_rows = [
        ("Document role", "Architecture, governance, purposes, assumptions, and review guidance"),
        ("Technical inventory", "AppSheet_Data_Dictionary.xlsx"),
        ("Source baseline", f"{data['sourceFile']} - configuration dated {data['sourceBaselineDate']}"),
        ("Conversion date", "30 August 2026 (Asia/Singapore)"),
        ("Reconciliation status", "Pending user reconciliation with live AppSheet configuration"),
        ("Maintenance model", "Excel is authoritative for row-level technical inventory; Word explains context and governance"),
    ]
    base.make_table(doc, ["Control", "Value"], [2700, 11556], cover_rows, font_size=9.2)
    base.add_para(doc, "", after=8)
    base.add_callout(doc, "Conversion boundary", "This split reproduces the existing Word baseline. It does not assert that the baseline still matches the live AppSheet database or either app as of the conversion date.")
    doc.add_page_break()

    doc.add_heading("1. Purpose, authority, and maintenance", level=1)
    base.add_para(doc, "This overview explains how the shared database, the OLTP Operations app, and the OLAP Reporting app fit together. Detailed tables, columns, expressions, application variances, slices, and slice-column membership are maintained in the companion Excel workbook.")
    doc.add_heading("1.1 Authority and status", level=2)
    base.add_label_value(doc, "System authority", "The AppSheet database and live app editors remain authoritative.")
    base.add_label_value(doc, "Conversion authority", "The content in this pair was converted from the existing Word baseline only.")
    base.add_label_value(doc, "Reconciliation status", "Pending user reconciliation; no AppSheet setup was opened or compared during this conversion.")
    base.add_label_value(doc, "Sensitive data", "The reference describes configuration only and contains no production records, borrower documents, bank details, or credentials.")

    doc.add_heading("1.2 Division of responsibility", level=2)
    responsibility_rows = [
        ("Word overview", "Narrative", "Architecture, ownership, purposes, assumptions, risks, and maintenance workflow", "Update when meaning, ownership, architecture, or governance changes"),
        ("Excel dictionary", "Technical inventory", "Tables, columns, expressions, relationships, slices, slice columns, and app variances", "Update for every data-model or slice change"),
        ("Excel change log", "Audit trail", "What changed, when, why, and whether it was reconciled with AppSheet", "Append one row for every accepted model revision"),
    ]
    base.make_table(doc, ["Artifact", "Role", "Contents", "Update trigger"], [2300, 1900, 5350, 4706], responsibility_rows, font_size=8.5)

    doc.add_heading("1.3 Maintenance procedure", level=2)
    add_numbered_steps(doc, [
        ("Inspect", "Review the affected source table and both app definitions before changing the reference."),
        ("Update Excel", "Edit the affected inventory rows first and preserve stable object names and column order."),
        ("Update Word", "Revise this overview only when architecture, purpose, ownership, assumptions, risk, or validation guidance changes."),
        ("Record", "Append a dated Change Log row in the workbook, including reconciliation status and source baseline."),
        ("Verify", "Check filters, formulas as text, relationships, slice membership, and cross-document consistency before accepting the revision."),
    ])
    base.add_callout(doc, "Maintenance rule", "Do not duplicate the full technical dictionary in Word. Excel is the row-level source of truth for this reference pair; Word supplies context.", fill="FFF8E8", color=base.CAUTION)

    doc.add_page_break()
    doc.add_heading("2. Architecture overview", level=1)
    copy_matrix_table(doc, data["architecture"], [2100, 3450, 2350, 6356], font_size=8.7)
    base.add_para(doc, "The OLTP and OLAP apps share the same Money database. Application-specific variance is primarily expressed through loaded tables, virtual columns, and curated slices. The Calendar helper remains an external Google Sheets-backed input to the OLAP model.")

    doc.add_heading("3. Table catalog", level=1)
    base.add_para(doc, "The companion workbook contains the complete column-level dictionary. This table remains in Word because source ownership and table purpose are architectural facts.")
    copy_matrix_table(doc, data["tables"], [1900, 2050, 1450, 900, 900, 900, 6156], font_size=8.0)
    base.add_para(doc, "Physical count includes the AppSheet or Sheets row ordinal where exposed. Row ID is the documented stable key for AppSheet Database tables; Calendar uses Date as its key.", size=8.5, italic=True, color=base.MUTED)

    doc.add_heading("4. Relationship model", level=1)
    base.add_para(doc, "The detailed workbook records the same direct Ref links in a filterable table. Reverse-reference virtual columns remain part of the Columns sheet.")
    copy_matrix_table(doc, data["relationships"], [2400, 2300, 2200, 3600, 3756], font_size=8.5)
    base.add_callout(doc, "Cascade caution", "Rows marked Is a part of? may follow the referenced parent lifecycle. Financially material corrections should remain auditable and should not rely on routine physical deletion.", fill="FFF1F0", color=base.RED)

    doc.add_heading("5. Slice catalog", level=1)
    slice_headers = ["App", "Slice", "Source", "Row filter", "Mode", "Cols", "Purpose"]
    slice_rows = [[s["App"], s["Slice"], s["Source"], s["Row Filter"], s["Update Mode"], s["Column Count"], s["Purpose"]] for s in data["slices"]]
    base.make_table(doc, slice_headers, [850, 2450, 1650, 3600, 1250, 650, 3806], slice_rows, font_size=7.7)
    base.add_para(doc, "Exact slice-column membership and order are maintained only in the workbook's Slice Columns sheet.", size=8.7, italic=True, color=base.MUTED)

    doc.add_page_break()
    doc.add_heading("6. Financial-model assumptions retained from the baseline", level=1)
    base.add_para(doc, "The following notes are carried forward verbatim in substance from the existing Word baseline. They have not been re-tested during this conversion.")
    for item in data["monthlyForecastAssumptions"]:
        if item.startswith("The existing KPI names"):
            base.add_para(doc, item)
        else:
            add_bullet(doc, item)
    base.add_callout(doc, "Validation status", "Treat these assumptions as pending confirmation until the user reconciles the converted reference with the current AppSheet formulas, bots, and business rules.", fill="FFF8E8", color=base.CAUTION)

    doc.add_heading("7. Known implications and follow-up checks", level=1)
    copy_matrix_table(doc, data["implications"], [2700, 11556], font_size=9.0)

    doc.add_heading("7.1 Manual validation checklist", level=2)
    for item in data["manualChecklist"]:
        label, value = parse_label(item)
        base.add_label_value(doc, label or "Check", value, after=4)

    doc.add_heading("8. Companion workbook map", level=1)
    workbook_rows = [
        ("README", "Status, maintenance rules, scope, and formula-driven inventory counts"),
        ("Tables", "One row per source table with app loading, physical/VC counts, and purpose"),
        ("Columns", "One row per unioned column definition, including type, availability, flags, purpose, and expression"),
        ("Relationships", "One row per direct Ref relationship and part-of behavior"),
        ("Slices", "One row per OLTP/OLAP slice with filter, update mode, purpose, and column count"),
        ("Slice Columns", "One row per included slice column, preserving display order"),
        ("App Variances", "Tables and columns present in only one app"),
        ("Change Log", "Append-only maintenance history and reconciliation status"),
    ]
    base.make_table(doc, ["Worksheet", "Purpose"], [2600, 11656], workbook_rows, font_size=9.0)
    base.add_callout(doc, "Next action", "Reconcile the workbook against the live AppSheet database and app editors, then update the Change Log status from Pending to Reconciled.")

    for table in doc.tables:
        set_repeat_header(table.rows[0])
        for row in table.rows:
            prevent_row_split(row)
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True

    doc.save(OUTPUT_DOCX)


def main():
    source_doc = Document(SOURCE_DOCX)
    inventory = json.loads(INVENTORY_JSON.read_text(encoding="utf-8"))
    data = make_conversion_data(source_doc, inventory)
    CONVERSION_JSON.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    build_overview(data)
    print(CONVERSION_JSON)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
