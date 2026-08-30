import json
from pathlib import Path
from docx import Document

root = Path(r"D:\GitHub\AppSheet Project")
data = json.loads((root / ".tmp" / "model_inventory.json").read_text(encoding="utf-8"))
print("json top keys:", list(data))
for key, value in data.items():
    if isinstance(value, dict):
        print(key, "dict keys:", list(value)[:20])
    elif isinstance(value, list):
        print(key, "list length:", len(value), "sample keys:", list(value[0]) if value else [])
    else:
        print(key, type(value).__name__, value)

for app_key, app in data["apps"].items():
    print("app", app_key, "keys", list(app))
    print(" tables", len(app.get("tables", [])), "slices", len(app.get("slices", [])))
    if app.get("tables"):
        t = app["tables"][0]
        print(" first table keys", list(t), "first row keys", list(t.get("rows", [{}])[0]))
    if app.get("slices"):
        print(" first slice", app["slices"][0])

print("relationships", data["relationships"])

doc = Document(root / "Documents" / "AppSheet_Data_Model_Reference.docx")
print("doc paragraphs:", len(doc.paragraphs), "tables:", len(doc.tables), "sections:", len(doc.sections))
for i, table in enumerate(doc.tables):
    head = [c.text.replace("\n", " | ")[:80] for c in table.rows[0].cells]
    print(f"table {i}: rows={len(table.rows)} cols={len(table.columns)} head={head}")
print("headings:")
for i, p in enumerate(doc.paragraphs):
    if p.style and p.style.name.startswith("Heading"):
        print(i, p.style.name, repr(p.text))

print("paragraph inventory:")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(i, p.style.name if p.style else "", repr(p.text))
